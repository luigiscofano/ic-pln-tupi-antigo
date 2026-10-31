"""Executa o núcleo do notebook recuperado com uma planilha sintética temporária.

Não chama upload/download do Colab nem usa dados lexicográficos reais.
"""
from contextlib import redirect_stdout
from copy import copy
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
import csv
import json
import os
import re
import unicodedata

import openpyxl
from openpyxl import Workbook

root = Path(__file__).resolve().parents[1]
notebook = json.loads((root / "acervo/tupi_carvalho_navarro_colab.ipynb").read_text(encoding="utf-8"))
codes = ["".join(c["source"]) for c in notebook["cells"] if c["cell_type"] == "code"]
regras = next(c for c in codes if "def normalize_navarro(" in c)
transformacao = next(c for c in codes if "wb_src = openpyxl.load_workbook" in c)
anterior = Path.cwd()
with TemporaryDirectory() as folder:
    try:
        os.chdir(folder)
        entrada = Path("entrada.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Fixture"
        ws.append(["TUPI", "PT_BR", "EXPLICACAO_PT_BR", "EN_US"])
        # Os textos são fixtures para testar o programa; não são verbetes validados.
        linhas = [("AI", "Fixture 1", "Texto, com vírgula", "Synthetic 1"),
                  ("EXEMPLO-Ã", "Fixture 2", 'Aspas " e acento', "Synthetic 2")]
        for linha in linhas:
            ws.append(linha)
        wb.save(entrada)
        contexto = {"re": re, "csv": csv, "unicodedata": unicodedata,
                    "copy": copy, "openpyxl": openpyxl, "Workbook": Workbook,
                    "input_path": str(entrada)}
        log = StringIO()
        with redirect_stdout(log):
            exec(compile(regras, "colab:regras", "exec"), contexto)
            exec(compile(transformacao, "colab:exportacao", "exec"), contexto)
        saida = openpyxl.load_workbook("tupi_carvalho_navarro_tabela.xlsx").active
        assert saida.max_row == 3 and saida.max_column == 5
        expected = ["TUPI O.O.", "TUPI O.N.", "PTBR", "ENUS", "EXPLICAÇÃO PTBR"]
        assert [c.value for c in saida[1]] == expected
        for i, linha in enumerate(linhas, 2):
            assert saida.cell(i, 1).value == linha[0]
            assert saida.cell(i, 2).value == contexto["normalize_navarro"](linha[0])
            assert [saida.cell(i, c).value for c in (3, 4, 5)] == [linha[1], linha[3], linha[2]]
        with open("tupi_carvalho_navarro_tabela.csv", encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file))
        assert rows == [[cell.value for cell in row] for row in saida.iter_rows()]
        assert saida.freeze_panes == "A2" and saida.auto_filter.ref == "A1:E3"
    finally:
        os.chdir(anterior)
print("PASS: núcleo do Colab recuperado; XLSX e CSV em fixture sintética")
